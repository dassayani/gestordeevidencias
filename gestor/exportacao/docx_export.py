"""Exportação em DOCX (Word), com os mesmos 3 modelos do `pdf_export.py`.
DOCX é um formato de fluxo (não de páginas fixas como o PDF), então aqui a
paginação e as faixas de cor são aproximadas — não pixel a pixel iguais ao
PDF, mas seguindo a mesma estrutura e paleta."""
import os

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

TEXTO = RGBColor(0x17, 0x26, 0x2E)
MUTED = RGBColor(0x6B, 0x81, 0x90)
MUTED_CLARO = RGBColor(0x8E, 0xA2, 0xAD)

_TWIPS_POR_CM = 1440 / 2.54  # o Word mede largura de tabela em twips (1/1440")


def _opcoes_padrao(opcoes):
    o = dict(opcoes or {})
    o.setdefault("cor_destaque", "#0B7285")
    o.setdefault("fonte_legenda", "Arial")
    o.setdefault("numerar_passos", True)
    # faltavam aqui: a borda era escolhida na tela e simplesmente ignorada
    # na geração do DOCX (o PDF já tratava, em pdf_export._opcoes_padrao)
    o.setdefault("borda_ativada", False)
    o.setdefault("borda_cor", "#0B7285")
    return o


def _capa_dados(capa):
    capa = capa or {}
    return {
        "titulo": capa.get("titulo") or "Evidências de teste",
        "caso": capa.get("caso", ""),
        "autor": capa.get("autor", ""),
        "data": capa.get("data", ""),
        "ambiente": capa.get("ambiente", ""),
    }


def _hex_limpo(hexcolor, padrao="0B7285"):
    """Normaliza pra 6 dígitos hex sem '#'; devolve o padrão se vier inválido."""
    texto = (hexcolor or "").lstrip("#").strip()
    if len(texto) == 6:
        try:
            int(texto, 16)
            return texto.upper()
        except ValueError:
            pass
    return padrao


def _rgb(hexcolor):
    # tolerante a hex inválido, como pdf_export._cor — antes um valor
    # quebrado levantava ValueError e derrubava a exportação inteira
    texto = _hex_limpo(hexcolor)
    return RGBColor(int(texto[0:2], 16), int(texto[2:4], 16), int(texto[4:6], 16))


def _sombrear_celula(celula, hexcolor):
    tcPr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    # os três atributos são necessários: só com w:fill o Word ignora o
    # sombreamento e a célula sai sem cor nenhuma
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), _hex_limpo(hexcolor))
    tcPr.append(shd)


def _largura_fixa(tabela, larguras_cm):
    """Fixa a largura das colunas de uma tabela.

    Definir só `cell.width` não basta: sem `w:tblLayout` fixo o Word recalcula
    as colunas pelo conteúdo e pode encolher uma delas bem abaixo do pedido —
    era o que espremia a lateral do Relatório QA a ponto de quebrar as
    palavras no meio ("Relatóri / o de teste").
    """
    tabela.autofit = False
    tblPr = tabela._tbl.tblPr

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)

    largura_total = OxmlElement("w:tblW")
    largura_total.set(qn("w:type"), "dxa")
    largura_total.set(qn("w:w"), str(int(sum(larguras_cm) * _TWIPS_POR_CM)))
    tblPr.append(largura_total)

    for idx, cm in enumerate(larguras_cm):
        if idx >= len(tabela.columns):
            break
        tabela.columns[idx].width = Cm(cm)
        for celula in tabela.columns[idx].cells:
            celula.width = Cm(cm)


def _sem_bordas(tabela):
    tbl = tabela._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for nome in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{nome}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)


def _bordas_externas(tabela, hexcolor):
    """Só o contorno da tabela, sem as linhas internas (o quadro do cartão)."""
    tblPr = tabela._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for nome in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{nome}")
        if nome.startswith("inside"):
            el.set(qn("w:val"), "none")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "6")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), _hex_limpo(hexcolor))
        borders.append(el)
    tblPr.append(borders)


def _linha_inteira(linha):
    """Impede o Word de partir a linha entre duas páginas (w:cantSplit)."""
    trPr = linha._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


def _sem_espaco(paragrafo, antes=0, depois=0):
    paragrafo.paragraph_format.space_before = Pt(antes)
    paragrafo.paragraph_format.space_after = Pt(depois)


def _configurar_pagina(document, fonte):
    secao = document.sections[0]
    secao.page_width = Mm(210)
    secao.page_height = Mm(297)
    document.styles["Normal"].font.name = fonte
    document.styles["Normal"].font.size = Pt(10.5)


def _borda_paragrafo(paragrafo, hexcolor):
    """Desenha um quadro em volta do parágrafo (w:pBdr).

    É como a borda das imagens é feita no DOCX: o python-docx não expõe
    contorno de imagem, então o quadro vai no parágrafo que contém a figura.
    Mesma técnica de injeção de OOXML já usada em `_sem_bordas`.
    """
    pPr = paragrafo._p.get_or_add_pPr()
    bordas = OxmlElement("w:pBdr")
    for nome in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{nome}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")        # 8 oitavos de ponto = 1pt
        el.set(qn("w:space"), "4")     # respiro entre a borda e a imagem
        el.set(qn("w:color"), _hex_limpo(hexcolor))
        bordas.append(el)
    pPr.append(bordas)


def _faixa_colorida(document, hexcolor, altura_pt=6):
    """Faixa fina na cor de destaque, no topo da página."""
    tabela = document.add_table(rows=1, cols=1)
    _sem_bordas(tabela)
    celula = tabela.rows[0].cells[0]
    _sombrear_celula(celula, hexcolor)
    p = celula.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    for run in p.runs:
        run.font.size = Pt(altura_pt)
    if not p.runs:
        p.add_run("").font.size = Pt(altura_pt)
    return tabela


# Passos que saíram sem imagem na última exportação (ver pdf_export).
_sem_imagem = []


def _largura_que_cabe(caminho, largura_cm, altura_max_cm):
    """Largura em que a imagem cabe na caixa largura x altura_max, sem distorcer.

    Só com a largura fixa um print mais alto que largo passava da altura
    prevista e empurrava o cartão seguinte para outra página.
    """
    if not altura_max_cm:
        return largura_cm
    try:
        from PIL import Image
        with Image.open(caminho) as img:
            w, h = img.size
    except Exception:
        return largura_cm
    if not w or not h:
        return largura_cm
    return min(largura_cm, altura_max_cm * w / h)


def _imagem_com_largura(paragrafo, caminho, largura_cm, opcoes=None,
                        altura_max_cm=None):
    run = paragrafo.add_run()
    inserida = False
    if caminho and os.path.exists(caminho):
        largura = Cm(_largura_que_cabe(caminho, largura_cm, altura_max_cm))
        try:
            run.add_picture(caminho, width=largura)
            inserida = True
        except Exception:
            # formato que o python-docx não lê: converte para RGB e tenta de novo
            try:
                import io
                from PIL import Image
                buffer = io.BytesIO()
                with Image.open(caminho) as img:
                    img.convert("RGB").save(buffer, format="PNG")
                buffer.seek(0)
                run.add_picture(buffer, width=largura)
                inserida = True
            except Exception:
                pass
    if not inserida:
        # lacuna visível, e não um parágrafo vazio que ninguém nota
        nome = os.path.basename(caminho or "") or "(sem arquivo)"
        _sem_imagem.append(nome)
        run.text = "[imagem indisponível: %s]" % nome
        run.font.color.rgb = MUTED_CLARO
    o = opcoes or {}
    if o.get("borda_ativada"):
        _borda_paragrafo(paragrafo, o.get("borda_cor"))


def _campo(paragrafo, instrucao):
    """Campo do Word (PAGE, NUMPAGES) que ele mesmo calcula ao abrir."""
    def marca(tipo):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), tipo)
        return el

    run = paragrafo.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instrucao} "
    valor = OxmlElement("w:t")
    valor.text = "1"  # o que aparece até o Word atualizar o campo
    for el in (marca("begin"), instr, marca("separate"), valor, marca("end")):
        run._r.append(el)
    return run


def _rodape_simples(document, texto):
    secao = document.sections[0]
    rodape = secao.footer
    p = rodape.paragraphs[0] if rodape.paragraphs else rodape.add_paragraph()
    p.text = texto
    p.style = document.styles["Normal"]
    # "Página X de Y" alinhado à direita, como no rodapé do PDF
    largura_util = secao.page_width - secao.left_margin - secao.right_margin
    p.paragraph_format.tab_stops.add_tab_stop(
        largura_util, alignment=WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\tPágina ")
    _campo(p, "PAGE")
    p.add_run(" de ")
    _campo(p, "NUMPAGES")
    for run in p.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED_CLARO


# ---------------------------------------------------------------- 1f
def exportar_passo_a_passo(destino, capa, passos, opcoes=None):
    o = _opcoes_padrao(opcoes)
    c = _capa_dados(capa)
    cor = _rgb(o["cor_destaque"])
    fonte = o["fonte_legenda"]

    document = Document()
    _configurar_pagina(document, fonte)
    secao = document.sections[0]
    secao.left_margin = secao.right_margin = Cm(1)
    secao.top_margin = secao.bottom_margin = Cm(1.2)

    # Faixa colorida no topo, equivalente ao pdf.rect(0, 0, PAGE_W, 3) do PDF.
    # Sem ela este modelo não tinha nenhum bloco na cor escolhida, e a cor
    # "sumia" mesmo estando aplicada nos textos.
    _faixa_colorida(document, o["cor_destaque"])

    eyebrow = document.add_paragraph()
    r = eyebrow.add_run("MANUAL DE OPERAÇÃO")
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = cor

    titulo = document.add_paragraph()
    rt = titulo.add_run(c["titulo"])
    rt.font.size = Pt(22)
    rt.font.bold = True
    rt.font.color.rgb = TEXTO

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    # caso e ambiente entram aqui: eram preenchidos na tela e não saíam
    # em lugar nenhum deste modelo. Vazios são omitidos.
    campos = [v for v in (c["caso"], c["ambiente"], c["autor"], c["data"]) if v]
    rm = meta.add_run(" · ".join(campos))
    rm.font.size = Pt(9)
    rm.font.color.rgb = MUTED
    document.add_paragraph()

    for idx, passo in enumerate(passos):
        p = document.add_paragraph()
        if o["numerar_passos"]:
            rn = p.add_run(f"{idx + 1}.  ")
            rn.font.bold = True
            rn.font.color.rgb = cor
        rl = p.add_run(passo.get("legenda") or "Sem legenda")
        rl.font.bold = True
        rl.font.size = Pt(12)
        rl.font.color.rgb = TEXTO

        img_p = document.add_paragraph()
        _imagem_com_largura(img_p, passo.get("caminho"), 19, o)
        if idx < len(passos) - 1:
            document.add_paragraph()

    _rodape_simples(document, "Gestor de Evidências")
    document.save(destino)


# ---------------------------------------------------------------- 1g
def exportar_ficha_evidencia(destino, capa, passos, opcoes=None):
    o = _opcoes_padrao(opcoes)
    c = _capa_dados(capa)
    cor = _rgb(o["cor_destaque"])
    fonte = o["fonte_legenda"]

    document = Document()
    _configurar_pagina(document, fonte)
    # Mesma geometria do PDF: margem lateral de 2cm, 2 evidências por página
    # e o cabeçalho colorido repetido em cada uma. Antes o cartão era uma
    # célula única que o Word partia entre páginas (legenda numa, imagem na
    # outra) e a imagem de 17cm vazava da área útil de 15,9cm.
    secao = document.sections[0]
    secao.left_margin = secao.right_margin = Cm(2)
    # topo zerado para a faixa do cabeçalho encostar na borda da folha
    secao.top_margin = Cm(0)
    secao.header_distance = Cm(0)
    secao.bottom_margin = Cm(1.8)
    secao.footer_distance = Cm(1)
    largura = 17.0
    # teto de altura da imagem: com legenda de até 4 linhas, dois cartões e o
    # cabeçalho ainda cabem numa página A4
    altura_img = 8.0

    def cabecalho():
        # faixa de borda a borda no topo, como o rect(0, 0, PAGE_W, 32) do
        # PDF: a tabela tem a largura da folha e recua a margem esquerda. As
        # células ficam sem margem interna (no modo de compatibilidade do
        # template o Word desloca a tabela pela margem da 1ª célula, e a faixa
        # começava 2cm fora da folha); o texto se alinha aos cartões por recuo.
        header = document.add_table(rows=1, cols=2)
        _sem_bordas(header)
        _largura_fixa(header, (14.0, 21.0 - 14.0))
        tblPr = header._tbl.tblPr
        recuo = OxmlElement("w:tblInd")
        recuo.set(qn("w:w"), str(int(-2 * _TWIPS_POR_CM)))
        recuo.set(qn("w:type"), "dxa")
        tblPr.append(recuo)
        _linha_inteira(header.rows[0])
        header.rows[0].height = Cm(3.2)
        header.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        esq, dir_ = header.rows[0].cells
        for cel in (esq, dir_):
            _sombrear_celula(cel, o["cor_destaque"])
            cel.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            # o w:tcMar tem que vir depois do w:shd, senão o Word o ignora
            tcMar = OxmlElement("w:tcMar")
            for lado in ("left", "right"):
                el = OxmlElement(f"w:{lado}")
                el.set(qn("w:w"), "0")
                el.set(qn("w:type"), "dxa")
                tcMar.append(el)
            cel._tc.get_or_add_tcPr().vAlign.addprevious(tcMar)
        p1 = esq.paragraphs[0]
        _sem_espaco(p1, antes=8)
        p1.paragraph_format.left_indent = Cm(2)
        r1 = p1.add_run("REGISTRO DE EVIDÊNCIAS")
        r1.font.size = Pt(8)
        r1.font.color.rgb = RGBColor(0xE6, 0xF0, 0xF2)
        p2 = esq.add_paragraph()
        _sem_espaco(p2, antes=2, depois=8)
        p2.paragraph_format.left_indent = Cm(2)
        titulo_caso = c["titulo"] + (f" · {c['caso']}" if c["caso"] else "")
        r2 = p2.add_run(titulo_caso)
        r2.font.size = Pt(15)
        r2.font.bold = True
        r2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p3 = dir_.paragraphs[0]
        _sem_espaco(p3)
        p3.paragraph_format.right_indent = Cm(2)
        p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        meta = [v for v in (c["ambiente"], c["data"], c["autor"]) if v]
        r3 = p3.add_run("\n".join(meta))
        r3.font.size = Pt(8)
        r3.font.color.rgb = RGBColor(0xCF, 0xE8, 0xEE)

    def espaco(pt, quebra_pagina=False):
        p = document.add_paragraph()
        _sem_espaco(p)
        r = p.add_run()
        r.font.size = Pt(pt)
        if quebra_pagina:
            r.add_break(WD_BREAK.PAGE)

    def cartao(numero, passo):
        tabela = document.add_table(rows=2, cols=2)
        _bordas_externas(tabela, "E1E8EE")
        _largura_fixa(tabela, (2.2, largura - 2.2))
        for linha in tabela.rows:
            _linha_inteira(linha)

        selo, cel_leg = tabela.rows[0].cells
        p_ev = selo.paragraphs[0]
        _sem_espaco(p_ev, antes=4, depois=4)
        p_ev.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ev.paragraph_format.keep_with_next = True
        _sombrear_celula(selo, "E6F2F5")
        selo.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        r_ev = p_ev.add_run(f"EV-{numero:02d}")
        r_ev.font.bold = True
        r_ev.font.size = Pt(8)
        r_ev.font.color.rgb = cor

        p_leg = cel_leg.paragraphs[0]
        _sem_espaco(p_leg, antes=3, depois=6)
        # legenda e imagem nunca em páginas diferentes
        p_leg.paragraph_format.keep_with_next = True
        r_leg = p_leg.add_run(passo.get("legenda") or "Sem legenda")
        r_leg.font.bold = True
        r_leg.font.size = Pt(11)
        r_leg.font.color.rgb = TEXTO

        # altura exata na linha da imagem: editando o DOCX à mão (Enter,
        # trocar o print) o cartão mantém o tamanho, como no PDF; só a
        # linha da legenda cresce se o texto aumentar
        tabela.rows[1].height = Cm(altura_img + 0.6)
        tabela.rows[1].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        a, b = tabela.rows[1].cells
        cel_img = a.merge(b)
        cel_img.width = Cm(largura)
        p_img = cel_img.paragraphs[0]
        _sem_espaco(p_img)
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cel_img.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _imagem_com_largura(p_img, passo.get("caminho"), largura - 1, o,
                            altura_max_cm=altura_img)

    grupos = list(range(0, max(1, len(passos)), 2))
    for g, idx in enumerate(grupos):
        cabecalho()
        espaco(8)
        par = [j for j in (idx, idx + 1) if j < len(passos)]
        for k, j in enumerate(par):
            cartao(j + 1, passos[j])
            ultimo_do_grupo = k == len(par) - 1
            # a quebra vai no próprio parágrafo separador: um parágrafo extra
            # só para ela podia cair sozinho numa página em branco
            espaco(2 if ultimo_do_grupo else 8,
                   quebra_pagina=ultimo_do_grupo and g < len(grupos) - 1)

    _rodape_simples(document, "Gestor de Evidências")
    document.save(destino)


# ---------------------------------------------------------------- 1h
def exportar_relatorio_qa(destino, capa, passos, opcoes=None):
    o = _opcoes_padrao(opcoes)
    c = _capa_dados(capa)
    cor = _rgb(o["cor_destaque"])
    fonte = o["fonte_legenda"]

    document = Document()
    _configurar_pagina(document, fonte)
    # margens estreitas como no modelo "passo": a tabela lateral+conteúdo tem
    # 19cm e não caberia na área de texto das margens padrão (15,9cm), que é
    # justamente o que fazia o Word encolher as colunas por conta própria.
    secao = document.sections[0]
    secao.left_margin = secao.right_margin = Cm(1)
    secao.top_margin = secao.bottom_margin = Cm(1.2)

    resumo = document.add_table(rows=1, cols=2)
    _sem_bordas(resumo)
    _largura_fixa(resumo, (5.5, 13.5))
    cel_lateral, cel_principal = resumo.rows[0].cells
    _sombrear_celula(cel_lateral, "0F3B45")

    p = cel_lateral.paragraphs[0]
    r = p.add_run("Relatório de teste")
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p2 = cel_lateral.add_paragraph()
    # mesmo conjunto de campos da lateral no PDF, que agora inclui o caso
    contexto = [v for v in (c["titulo"], c["caso"], c["ambiente"]) if v]
    r2 = p2.add_run("\n".join(contexto))
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = RGBColor(0x9F, 0xC4, 0xCC)
    p3 = cel_lateral.add_paragraph()
    r3 = p3.add_run(f"\n{len(passos)}")
    r3.font.bold = True
    r3.font.size = Pt(20)
    r3.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p4 = cel_lateral.add_paragraph()
    r4 = p4.add_run("capturas neste relatório")
    r4.font.size = Pt(8)
    r4.font.color.rgb = RGBColor(0x9F, 0xC4, 0xCC)
    p5 = cel_lateral.add_paragraph()
    r5 = p5.add_run(f"\n{c['autor']}\n{c['data']}")
    r5.font.size = Pt(7.5)
    r5.font.color.rgb = RGBColor(0x9F, 0xC4, 0xCC)

    pp = cel_principal.paragraphs[0]
    rp = pp.add_run("EVIDÊNCIAS")
    rp.font.size = Pt(9)
    rp.font.bold = True
    rp.font.color.rgb = cor

    for idx, passo in enumerate(passos):
        pn = cel_principal.add_paragraph()
        rn = pn.add_run(f"{idx + 1}.  ")
        rn.font.bold = True
        rn.font.size = Pt(13)
        rn.font.color.rgb = cor
        rl = pn.add_run(passo.get("legenda") or "Sem legenda")
        rl.font.bold = True
        rl.font.size = Pt(10.5)
        rl.font.color.rgb = TEXTO
        pim = cel_principal.add_paragraph()
        _imagem_com_largura(pim, passo.get("caminho"), 13, o)

    _rodape_simples(document, "Gestor de Evidências")
    document.save(destino)


MODELOS = {
    "passo": exportar_passo_a_passo,
    "ficha": exportar_ficha_evidencia,
    "qa": exportar_relatorio_qa,
}


def exportar(modelo, destino, capa, passos, opcoes=None):
    """Gera o DOCX. Devolve os nomes dos passos que ficaram sem imagem."""
    funcao = MODELOS.get(modelo, MODELOS["passo"])
    del _sem_imagem[:]
    funcao(destino, capa, passos, opcoes)
    return list(_sem_imagem)
