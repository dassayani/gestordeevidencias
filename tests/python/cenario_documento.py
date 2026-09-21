# -*- coding: utf-8 -*-
"""Montar documento: campos, sequencia, previa e geracao nos tres modelos.

Percorre o caminho inteiro do usuario: escolhe varios prints, preenche a capa,
escreve a legenda de cada passo, reordena a sequencia, abre a previa e gera PDF
e DOCX. Depois le o arquivo gerado e confere que o conteudo corresponde ao que
foi configurado.

As caixas de dialogo (mensagem de sucesso, abrir a pasta no Explorer) sao
substituidas durante o teste: modal sem ninguem para clicar trava o processo.
"""

import os as _os
import sys as _sys

# Raiz do projeto a partir deste arquivo: tests/python/x.py -> duas pastas acima
RAIZ = _os.path.dirname(_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
_sys.path.insert(0, RAIZ)
_os.chdir(RAIZ)
SAIDA = _os.path.join(RAIZ, "tests", "_saida")
_os.makedirs(SAIDA, exist_ok=True)

import os, sys, ctypes, tempfile, shutil, traceback

if not ctypes.windll.user32.SetProcessDpiAwarenessContext(ctypes.c_void_p(-4)):
    ctypes.windll.shcore.SetProcessDpiAwareness(2)

from tkinterdnd2 import TkinterDnD
from PIL import Image
import pymupdf
import config
import document_builder
import export_preview
from workspace import AppEvidencias

falhas = []
avisos = []


def checar(condicao, descricao):
    print(("OK:     " if condicao else "FALHOU: ") + descricao)
    if not condicao:
        falhas.append(descricao)


class DialogoFalso:
    """Substitui o messagebox: guarda o que seria mostrado, sem travar."""

    @staticmethod
    def showinfo(titulo, mensagem, **kw):
        avisos.append(("info", titulo, mensagem))

    @staticmethod
    def showerror(titulo, mensagem, **kw):
        avisos.append(("erro", titulo, mensagem))

    @staticmethod
    def showwarning(titulo, mensagem, **kw):
        avisos.append(("aviso", titulo, mensagem))

    @staticmethod
    def askyesno(titulo, mensagem, **kw):
        avisos.append(("pergunta", titulo, mensagem))
        return True


export_preview.messagebox = DialogoFalso
export_preview.os.startfile = lambda caminho: avisos.append(("pasta", caminho, ""))

CAPA = {"titulo": "Validacao do fechamento mensal",
        "caso": "CT-7788",
        "autor": "QA - Dasayani",
        "data": "21/09/2026",
        "ambiente": "Homologacao"}
LEGENDAS = ["Tela inicial do modulo de fechamento",
            "Confirmacao do saldo disponivel na competencia",
            "Mensagem de sucesso apresentada ao usuario"]

tmp = tempfile.mkdtemp(prefix="ge_doc_")
capturas = os.path.join(tmp, "Capturas")
os.makedirs(capturas)
nomes = []
for i in range(3):
    nome = "print_30000%d.png" % i
    img = Image.new("RGB", (900, 560), (40 + i * 50, 100, 150))
    img.save(os.path.join(capturas, nome))
    nomes.append(nome)

cfg = config.load(tmp)
cfg["pasta_capturas"] = capturas
config.save(tmp, cfg)

root = TkinterDnD.Tk()
app = AppEvidencias(root, capturas, os.path.join(tmp, "PDF"), RAIZ, tmp)
app.pausar_timer = True
root.deiconify()
root.update()


def texto_do_documento(caminho):
    """Todo o texto do documento, com os espacos normalizados.

    Normalizar importa: no relatorio QA o titulo e quebrado em varias linhas
    pela largura da faixa lateral, e sem isto ele nunca bateria como uma
    string continua. No DOCX, a Ficha e o QA montam o conteudo em tabelas, que
    nao aparecem em `document.paragraphs`.
    """
    if caminho.lower().endswith(".pdf"):
        with pymupdf.open(caminho) as doc:
            bruto = "\n".join(p.get_text() for p in doc)
    else:
        import docx
        documento = docx.Document(caminho)
        partes = [p.text for p in documento.paragraphs]
        for tabela in documento.tables:
            for linha in tabela.rows:
                for celula in linha.cells:
                    partes.extend(p.text for p in celula.paragraphs)
        bruto = "\n".join(partes)
    return " ".join(bruto.split())


def normalizar(texto):
    return " ".join(str(texto).split())


try:
    janela = document_builder.MontarDocumento(app, nomes)
    janela.update()
    checar(len(janela.passos) == 3, "os 3 prints entraram no documento")

    # ---- preencher todos os campos da capa ----
    for chave, valor in CAPA.items():
        if chave in janela.entries_capa:
            entry = janela.entries_capa[chave]
            entry.delete(0, "end")
            entry.insert(0, valor)
    janela.update()
    lidos = {c: e.get().strip() for c, e in janela.entries_capa.items()}
    for chave, valor in CAPA.items():
        checar(lidos.get(chave) == valor, "o campo '%s' da capa aceitou o valor" % chave)

    # ---- legenda de cada passo ----
    for idx, legenda in enumerate(LEGENDAS):
        widget = janela.entries_legenda[idx]
        widget.delete("1.0", "end")
        widget.insert("1.0", legenda)
        janela._on_legenda_change(idx, widget)
    janela.update()
    checar([p["legenda"] for p in janela.passos] == LEGENDAS,
           "as legendas dos passos foram guardadas")

    # ---- reordenar a sequencia ----
    primeiro = janela.passos[0]["legenda"]
    janela._mover(0, 1)
    janela.update()
    checar(janela.passos[1]["legenda"] == primeiro,
           "mover para baixo troca a ordem dos passos")
    janela._mover(1, -1)
    janela.update()
    checar(janela.passos[0]["legenda"] == primeiro, "mover de volta restaura a ordem")

    # ---- gerar nos tres modelos, nos dois formatos ----
    for modelo in ("passo", "ficha", "qa"):
        janela.var_modelo.set(modelo)
        janela.var_cor.set("#E8590C")
        janela.var_borda.set(True)
        janela.update()

        janela._pre_visualizar()
        root.update()
        previa = None
        for w in root.winfo_children():
            if isinstance(w, export_preview.PreVisualizarExportar):
                previa = w
        checar(previa is not None, "[%s] a previa abriu" % modelo)
        if previa is None:
            continue

        checar(previa.modelo == modelo, "[%s] a previa usa o modelo escolhido" % modelo)
        checar(previa.capa.get("titulo") == CAPA["titulo"],
               "[%s] a previa recebeu a capa preenchida" % modelo)
        checar(previa.opcoes.get("cor_destaque") == "#E8590C",
               "[%s] a previa recebeu a cor escolhida" % modelo)
        checar(previa.opcoes.get("borda_ativada") is True,
               "[%s] a previa recebeu a opcao de borda" % modelo)

        destino = os.path.join(tmp, "saida_%s" % modelo)
        os.makedirs(destino, exist_ok=True)
        for formato in ("pdf", "docx"):
            previa.var_formato.set(formato)
            previa._entry_destino.config(state="normal")
            previa._entry_destino.delete(0, "end")
            previa._entry_destino.insert(0, destino)
            previa._exportar()
            root.update()
            gerados = [f for f in os.listdir(destino) if f.lower().endswith(formato)]
            checar(bool(gerados), "[%s] gerou o %s" % (modelo, formato.upper()))
            if gerados:
                caminho = os.path.join(destino, gerados[0])
                conteudo = texto_do_documento(caminho)
                checar(normalizar(CAPA["titulo"]) in conteudo,
                       "[%s/%s] o documento traz o titulo configurado" % (modelo, formato))
                checar(normalizar(CAPA["caso"]) in conteudo,
                       "[%s/%s] o documento traz o caso configurado" % (modelo, formato))
                checar(normalizar(LEGENDAS[0])[:30] in conteudo,
                       "[%s/%s] o documento traz a legenda do primeiro passo"
                       % (modelo, formato))
                checar("Gestor de Evidências" in conteudo or formato == "docx",
                       "[%s/%s] o documento traz o rodape padrao" % (modelo, formato))

            checar(bool(previa.winfo_exists()),
                   "[%s/%s] a previa continua aberta depois de gerar" % (modelo, formato))

        previa.destroy()
        root.update()

    checar(not [a for a in avisos if a[0] == "erro"],
           "nenhuma mensagem de erro durante as geracoes")
    janela.destroy()
    root.update()
except Exception:
    falhas.append("erro: " + traceback.format_exc())

print("\nFALHAS:", "nenhuma" if not falhas else "")
for f in falhas:
    print(" -", f)
try:
    root.destroy()
except Exception:
    pass
shutil.rmtree(tmp, ignore_errors=True)
sys.stdout.flush()
os._exit(0)
