# -*- coding: utf-8 -*-
"""Protecao contra perda de dados e contra falha calada (sem janela).

Defeitos encontrados na revisao critica do repositorio:
  * com o padrao de nome `hora` (o de fabrica, so HHMMSS) uma captura nova
    sobrescrevia a de outro dia no mesmo horario, e herdava o .raw.png e as
    anotacoes da antiga
  * a retencao so poupava o que tinha `edited_at`: a legenda digitada ao
    montar um documento (gravada sem marcar edicao) nao protegia a captura
  * a galeria decodificava todos os PNG a cada atualizacao
  * um config.json ilegivel virava "tudo padrao" em silencio, e a proxima
    gravacao apagava o que havia
  * o documento saia com o passo em branco quando a imagem nao entrava

Cobre cada um, mais o PNG de 16 bits que o fpdf 1.7 recusava.
"""

import os as _os
import sys as _sys

RAIZ = _os.path.dirname(_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
_sys.path.insert(0, RAIZ)
_os.chdir(RAIZ)

import os, sys, tempfile, shutil, time, json, warnings
from datetime import datetime

import numpy as np
import PIL.Image
from PIL import Image
from gestor.captura import captura_utils
from gestor.dados import capture_store as cs
from gestor.dados import config
from gestor.exportacao import docx_export, pdf_export

warnings.simplefilter("ignore", DeprecationWarning)
falhas = []


def checar(condicao, descricao):
    print(("OK: " if condicao else "FALHOU: ") + descricao)
    if not condicao:
        falhas.append(descricao)


tmp = tempfile.mkdtemp(prefix="ge_protecao_")


def png(nome, cor=(10, 120, 160), tamanho=(64, 48), pasta=None):
    caminho = os.path.join(pasta or tmp, nome)
    Image.new("RGB", tamanho, cor).save(caminho)
    return caminho


try:
    # ================= nomes que nao colidem =================
    pasta = os.path.join(tmp, "nomes")
    os.makedirs(pasta)
    agora = datetime(2026, 9, 21, 14, 30, 52)
    checar(captura_utils.nome_livre(pasta, "hora", agora) == "print_143052.png",
           "pasta vazia: o nome e o do padrao")
    png("print_143052.png", pasta=pasta)
    checar(captura_utils.nome_livre(pasta, "hora", agora) == "print_143052_2.png",
           "ja existe um print com esse horario (de outro dia): ganha sufixo, nao sobrescreve")
    png("print_143052_2.png", pasta=pasta)
    checar(captura_utils.nome_livre(pasta, "hora", agora) == "print_143052_3.png",
           "e o sufixo continua contando")
    with open(os.path.join(pasta, "print_101010.json"), "w") as f:
        f.write("{}")
    checar(captura_utils.nome_livre(pasta, "hora", datetime(2026, 1, 1, 10, 10, 10))
           == "print_101010_2.png",
           "um .json orfao tambem ocupa o nome (senao a captura nova herdaria os metadados)")
    png("print_202020.raw.png", pasta=pasta)
    checar(captura_utils.nome_livre(pasta, "hora", datetime(2026, 1, 1, 20, 20, 20))
           == "print_202020_2.png",
           "um .raw.png orfao tambem (senao o editor abriria a imagem errada)")
    png("print_20260921_143052.png", pasta=pasta)
    checar(captura_utils.nome_livre(pasta, "data_hora", agora) == "print_20260921_143052_2.png",
           "o padrao data_hora tambem ganha sufixo no mesmo segundo")
    with_sufixo = os.path.join(pasta, "print_20260921_143052_2.png")
    checar(cs.instante_da_captura(with_sufixo, {}) == agora.timestamp(),
           "o nome com sufixo ainda informa a data e a hora")
    so_hora = png("print_093000_2.png", pasta=pasta)
    stamp = datetime(2026, 9, 21, 16, 0, 0).timestamp()
    os.utime(so_hora, (stamp, stamp))
    checar(cs.instante_da_captura(so_hora, {"edited_at": "x"}) == datetime(2026, 9, 21, 9, 30).timestamp(),
           "e o nome so de hora com sufixo tambem")

    # ================= retencao =================
    pasta = os.path.join(tmp, "retencao")
    os.makedirs(pasta)
    velho = time.time() - 40 * 86400
    casos = {}
    for nome, meta in (("sem_nada", None),
                       ("com_legenda", {"caption": "Tela de login", "shapes": []}),
                       ("com_caso", {"caso": "CT-1", "shapes": []}),
                       ("com_anotacao", {"shapes": [{"id": 1}]}),
                       ("editada", {"edited_at": "2026-01-01T00:00:00", "shapes": []})):
        caminho = png("print_%s.png" % nome, pasta=pasta)
        if meta is not None:
            cs.save_meta(caminho, meta, marcar_editada=False)
        os.utime(caminho, (velho, velho))
        casos[nome] = caminho
    recente = png("print_recente.png", pasta=pasta)
    apagados = []
    original_delete = cs.delete_capture
    cs.delete_capture = lambda caminho: (apagados.append(os.path.basename(caminho)), True)[1]
    try:
        n = captura_utils.limpar_capturas_antigas(pasta, 30)
    finally:
        cs.delete_capture = original_delete
    checar(apagados == ["print_sem_nada.png"] and n == 1,
           "retencao apaga so a captura antiga que ninguem usou (apagou %s)" % apagados)
    checar("print_com_legenda.png" not in apagados,
           "captura com legenda (como a digitada no Montar) nao e apagada")
    checar("print_com_caso.png" not in apagados and "print_com_anotacao.png" not in apagados,
           "nem com caso/projeto, nem com anotacao")
    checar(os.path.basename(recente) not in apagados, "nem a recente")

    # ================= cache de miniaturas =================
    pasta = os.path.join(tmp, "miniaturas")
    os.makedirs(pasta)
    grande = png("print_cache.png", tamanho=(1920, 1080), pasta=pasta)
    abertos = []
    original_open = PIL.Image.open
    PIL.Image.open = lambda *a, **k: (abertos.append(a[0]), original_open(*a, **k))[1]
    try:
        a = cs.miniatura(grande, (116, 72))
        b = cs.miniatura(grande, (116, 72))
        primeira = len(abertos)
        derivada = cs.miniatura(grande, (60, 40), origem=(116, 72))
        depois_derivada = len(abertos)
        time.sleep(0.05)
        Image.new("RGB", (1920, 1080), (200, 0, 0)).save(grande)        # "editou"
        c = cs.miniatura(grande, (116, 72))
    finally:
        PIL.Image.open = original_open
    checar(primeira == 1 and a is b, "a segunda miniatura da mesma imagem vem do cache, sem abrir o PNG")
    checar(depois_derivada == 1 and derivada.size[0] <= 60,
           "a medida derivada sai da que ja esta em cache, sem decodificar de novo")
    checar(len(abertos) == 2 and c is not a and c.getpixel((5, 5))[0] > 150,
           "depois de editar o arquivo, a miniatura e refeita (a edicao aparece)")
    checar(a.size[0] <= 116 and a.size[1] <= 72 and a.mode == "RGB", "no tamanho e modo pedidos")
    limite = cs._LIMITE_MINIATURAS
    cs._LIMITE_MINIATURAS = 3
    try:
        for i in range(5):
            cs.miniatura(png("print_lru_%d.png" % i, pasta=pasta), (10, 10))
        checar(len(cs._MINIATURAS) <= 3, "o cache respeita o limite de memoria (%d)" % len(cs._MINIATURAS))
    finally:
        cs._LIMITE_MINIATURAS = limite

    # ================= config.json =================
    pasta = os.path.join(tmp, "config")
    os.makedirs(pasta)
    cfg = config.load(pasta)
    cfg["pasta_capturas"] = r"D:\Evidencias\Projeto"
    checar(config.save(pasta, cfg) is True, "salvar devolve True")
    checar(not [n for n in os.listdir(pasta) if n.endswith(".tmp")],
           "salvar nao deixa temporario")
    checar(config.load(pasta)["pasta_capturas"] == r"D:\Evidencias\Projeto", "e o valor volta")
    with open(os.path.join(pasta, "config.json"), "w", encoding="utf-8") as f:
        f.write('{"pasta_capturas": "D:\\\\Evidencias", "autor_pa')          # cortado no meio
    lido = config.load(pasta)
    guardados = [n for n in os.listdir(pasta) if ".ilegivel-" in n]
    checar(lido["pasta_capturas"] is None, "config.json ilegivel: o app abre com os padroes")
    checar(len(guardados) == 1, "mas o arquivo ilegivel e guardado ao lado, e nao perdido")
    if guardados:
        with open(os.path.join(pasta, guardados[0]), encoding="utf-8") as f:
            checar("Evidencias" in f.read(), "com o conteudo original, para recuperar a mao")
    config.save(pasta, lido)
    checar(len([n for n in os.listdir(pasta) if ".ilegivel-" in n]) == 1,
           "a gravacao seguinte nao apaga a copia guardada")
    with open(os.path.join(pasta, "config.json"), "w", encoding="utf-8") as f:
        json.dump(["nao", "e", "objeto"], f)
    checar(config.load(pasta)["pasta_capturas"] is None, "JSON que nao e objeto tambem cai nos padroes")
    checar(len([n for n in os.listdir(pasta) if ".ilegivel-" in n]) >= 1,
           "e tambem e guardado")

    # ================= imagem que nao entra no documento =================
    pasta = os.path.join(tmp, "export")
    os.makedirs(pasta)
    bom = png("print_bom.png", pasta=pasta)
    corrompido = os.path.join(pasta, "print_corrompido.png")
    with open(corrompido, "wb") as f:
        f.write(b"nao e png")
    ausente = os.path.join(pasta, "print_ausente.png")
    dezesseis = os.path.join(pasta, "print_16bits.png")
    Image.fromarray((np.ones((48, 64)) * 40000).astype(np.uint16)).save(dezesseis)
    passos = [{"caminho": c, "legenda": os.path.basename(c)}
              for c in (bom, corrompido, ausente, dezesseis)]
    import pymupdf
    import docx as _docx
    for modelo in ("passo", "ficha", "qa"):
        destino = os.path.join(pasta, "doc_%s.pdf" % modelo)
        sem = pdf_export.exportar(modelo, destino, {"titulo": "T"}, passos, {})
        checar(sorted(sem) == ["print_ausente.png", "print_corrompido.png"],
               "[pdf/%s] a exportacao informa os passos sem imagem (%s)" % (modelo, sem))
        with pymupdf.open(destino) as doc:
            # imagens unicas: o fpdf 1.7 poe todas num dicionario de recursos
            # compartilhado, e contar por pagina contaria em dobro
            imagens = len({img[0] for p in doc for img in p.get_images()})
            texto = " ".join(p.get_text() for p in doc)
        checar(imagens == 2, "[pdf/%s] o PNG de 16 bits entra (2 imagens, nao 1)" % modelo)
        checar(texto.count("imagem indisponível") == 2,
               "[pdf/%s] o lugar de cada imagem que faltou aparece marcado" % modelo)
        checar("print_bom.png" in texto and "print_16bits.png" in texto,
               "[pdf/%s] e as legendas continuam no documento" % modelo)
        destino = os.path.join(pasta, "doc_%s.docx" % modelo)
        sem = docx_export.exportar(modelo, destino, {"titulo": "T"}, passos, {})
        checar(sorted(sem) == ["print_ausente.png", "print_corrompido.png"],
               "[docx/%s] idem no DOCX (%s)" % (modelo, sem))
        documento = _docx.Document(destino)
        textos = [p.text for p in documento.paragraphs]
        for tabela in documento.tables:
            for linha in tabela.rows:
                for celula in linha.cells:
                    textos.extend(p.text for p in celula.paragraphs)
        checar(sum("imagem indisponível" in t for t in textos) == 2,
               "[docx/%s] com a lacuna marcada no texto" % modelo)
    checar(pdf_export.exportar("passo", os.path.join(pasta, "ok.pdf"), {}, passos[:1], {}) == [],
           "sem problema, a lista volta vazia (e nao herda a exportacao anterior)")
except Exception:
    import traceback
    falhas.append("erro: " + traceback.format_exc())

shutil.rmtree(tmp, ignore_errors=True)
print()
print("FALHAS:", falhas if falhas else "nenhuma")
sys.exit(1 if falhas else 0)
